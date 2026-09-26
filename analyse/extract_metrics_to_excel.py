"""把 outputs/ 里所有 train_log.txt 的最终 TEST RESULTS 汇总成一个 Excel。

用法：
    python analyse/extract_metrics_to_excel.py

输出：docs/experiment_metrics.xlsx（每次运行覆盖更新）。
只读取每个 train_log.txt 最后一个完整的 `=== TEST RESULTS === ... === END TEST RESULTS ===`
区块的正式指标，不用 val-best 行或 checkpoint 文件名。

表结构：
    Tag | Run | Experiment | Dataset | Recall | Precision | OA | F1 | IoU | Kappa
    | Params(M) | Trainable(M) | FLOPs(G) | ReparamErr | RepMode
- Tag/Run/Experiment/Dataset 由 outputs/ 下相对路径拆出，整表开启自动筛选，
  可按数据集或实验单独筛选。
- 六项指标统一 ×100 并保留两位小数（0.9877 -> 98.77）。
"""
import os
import re
import glob

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUTS = os.path.join(ROOT, "outputs")
DOCS = os.path.join(ROOT, "docs")
OUT_XLSX = os.path.join(DOCS, "experiment_metrics.xlsx")

# 六项指标（均为 0~1 比例，输出时 ×100）
METRIC_KEYS = ["Recall", "Precision", "OA", "F1", "IoU", "Kappa"]


def extract_test_block(text):
    """返回最后一个 TEST RESULTS 区块文本，没有则 None。"""
    blocks = re.findall(r"=== TEST RESULTS ===\n(.*?)\n=== END TEST RESULTS ===", text, re.S)
    return blocks[-1] if blocks else None


def parse_block(block):
    """解析一个 TEST RESULTS 区块，返回 dict。"""
    d = {}
    # 六项指标行： Recall=x | Precision=y | OA=z | F1=w | IoU=v | Kappa=u
    m = re.search(r"Recall=([0-9.]+)\s*\|\s*Precision=([0-9.]+)\s*\|\s*OA=([0-9.]+)\s*\|\s*F1=([0-9.]+)\s*\|\s*IoU=([0-9.]+)\s*\|\s*Kappa=([0-9.]+)", block)
    if m:
        d["Recall"], d["Precision"], d["OA"], d["F1"], d["IoU"], d["Kappa"] = [
            float(x) for x in m.groups()
        ]
    # 各种参数/FLOPs 行（baseline 与 TAR-DCR 两种格式）
    def _num(pattern):
        mm = re.search(pattern + r"\s+([0-9.]+)", block)
        return float(mm.group(1)) if mm else None

    d["Params(M)"] = _num(r"\[(?:PARAMS|TOTAL-TRAIN-GRAPH-PARAMS|DEPLOY-PARAMS)\]")
    d["Trainable(M)"] = _num(r"\[TRAINABLE-PARAMS\]")
    d["FLOPs(G)"] = _num(r"\[(?:FLOPS|DEPLOY-FLOPS)\]")
    err = re.search(r"\[REPARAM-MAX-ABS-ERROR\]\s+([0-9.eE+-]+)", block)
    d["ReparamErr"] = float(err.group(1)) if err else None
    rm = re.search(r"\[REP-MODE\]\s+(\w+)", block)
    d["RepMode"] = rm.group(1) if rm else None
    return d


def parse_path(rel):
    """把相对路径拆成 (tag, run, experiment, dataset)。

    例：
      baseline/Run1/WHU-CD-256            -> (baseline, Run1, "", WHU-CD-256)
      TAR-DCR/Run1/A0_Plain/CDD-CD-256    -> (TAR-DCR, Run1, A0_Plain, CDD-CD-256)
      TAR-DCR/Run2/full_last2/LEVIR-CD-256 -> (TAR-DCR, Run2, full_last2, LEVIR-CD-256)
      TAR-DCR/Run3/LEVIR-CD-256           -> (TAR-DCR, Run3, "", LEVIR-CD-256)
    """
    parts = rel.split("/")
    dataset = parts[-1]
    rest = parts[:-1]
    tag = rest[0] if rest else ""
    run = ""
    exp_parts = []
    for p in rest[1:]:
        if re.match(r"^Run\d+$", p):
            run = p
        else:
            exp_parts.append(p)
    experiment = "/".join(exp_parts)
    return tag, run, experiment, dataset


def main():
    logs = sorted(glob.glob(os.path.join(OUTPUTS, "**", "train_log.txt"), recursive=True))
    rows = []
    for log in logs:
        rel = os.path.relpath(os.path.dirname(log), OUTPUTS).replace("\\", "/")
        with open(log, encoding="utf-8", errors="replace") as f:
            text = f.read()
        block = extract_test_block(text)
        if block is None:
            print(f"[skip] {rel}: no TEST RESULTS block")
            continue
        info = parse_block(block)
        if "F1" not in info:
            print(f"[skip] {rel}: no metric line")
            continue
        rows.append((rel, info))
        print(f"[ok] {rel}: F1={info['F1'] * 100:.2f}")

    if not rows:
        print("没有找到任何 TEST RESULTS 区块。")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "metrics"
    headers = ["Tag", "Run", "Experiment", "Dataset"] + METRIC_KEYS + [
        "Params(M)", "Trainable(M)", "FLOPs(G)", "ReparamErr", "RepMode"
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for rel, info in rows:
        tag, run, experiment, dataset = parse_path(rel)
        metrics = [
            round(info[k] * 100, 2) if info.get(k) is not None else None
            for k in METRIC_KEYS
        ]
        row = [tag, run, experiment, dataset] + metrics + [
            info.get("Params(M)"), info.get("Trainable(M)"), info.get("FLOPs(G)"),
            info.get("ReparamErr"), info.get("RepMode"),
        ]
        ws.append(row)

    # 数值格式：六项指标两位小数；参数量/FLOPs 两位小数；折叠误差科学计数
    for r in range(2, ws.max_row + 1):
        for col in ["E", "F", "G", "H", "I", "J"]:
            ws[f"{col}{r}"].number_format = "0.00"
        ws[f"K{r}"].number_format = "0.00"      # Params(M)
        ws[f"L{r}"].number_format = "0.00"      # Trainable(M)
        ws[f"M{r}"].number_format = "0.00"      # FLOPs(G)
        ws[f"N{r}"].number_format = "0.00E+00"  # ReparamErr

    # 整表自动筛选（按数据集/实验/run/tag 均可筛）
    ws.auto_filter.ref = ws.dimensions

    # 冻结首行
    ws.freeze_panes = "A2"

    # 列宽
    widths = {
        "A": 10, "B": 8, "C": 14, "D": 16,
        "E": 10, "F": 10, "G": 8, "H": 8, "I": 8, "J": 8,
        "K": 10, "L": 12, "M": 10, "N": 12, "O": 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    os.makedirs(DOCS, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"\n已写入 {OUT_XLSX}（{len(rows)} 行）")


if __name__ == "__main__":
    main()
